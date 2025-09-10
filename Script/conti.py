import os
import re
import xlsxwriter
from PyPDF2 import PdfReader
from docx import Document

def count_word_in_pdf(file_path, word="continued"):
    try:
        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        return len(re.findall(rf"\b{word}\b", text, flags=re.IGNORECASE))
    except Exception as e:
        print(f"Error reading PDF {file_path}: {e}")
        return 0

from docx import Document

from docx import Document
import re

def count_word_in_docx(file_path, word="continued"):
    try:
        doc = Document(file_path)
        full_text = []

        # Only extract text from headers
        for section in doc.sections:
            if section.header:
                for para in section.header.paragraphs:
                    full_text.append(para.text)

        # Join all extracted header text and count occurrences of the word
        all_text = "\n".join(full_text)
        #print(all_text)
        return len(re.findall(rf"\b{word}\b", all_text, flags=re.IGNORECASE))

    except Exception as e:
        print(f"Error reading DOCX {file_path}: {e}")
        return 0



def match_and_count(folder_path, word="continued"):
    files = os.listdir(folder_path)
    pdf_files = {os.path.splitext(f)[0]: os.path.join(folder_path, f)
                 for f in files if f.lower().endswith('.pdf')}
    docx_files = {os.path.splitext(f)[0]: os.path.join(folder_path, f)
                  for f in files if f.lower().endswith('.docx')}

    results = []

    for pdf_base, pdf_path in pdf_files.items():
        # Try to find a matching DOCX that starts with the PDF name
        matching_docx = next((docx_path for docx_base, docx_path in docx_files.items()
                              if docx_base.startswith(pdf_base)), None)

        if matching_docx:
            pdf_count = count_word_in_pdf(pdf_path, word)
            docx_count = count_word_in_docx(matching_docx, word)
            results.append((pdf_base, pdf_count, docx_count))
        else:
            print(f"⚠️ No matching DOCX found for PDF: {pdf_base}")

    return results

def write_to_excel(results, output_file):
    workbook = xlsxwriter.Workbook(output_file)
    worksheet = workbook.add_worksheet("Continued Word Count")

    headers = ["Base Filename", "PDF Continued Count", "DOCX Continued Count"]
    for col_num, header in enumerate(headers):
        worksheet.write(0, col_num, header)

    for row_num, (filename, pdf_count, docx_count) in enumerate(results, start=1):
        worksheet.write(row_num, 0, filename)
        worksheet.write(row_num, 1, pdf_count)
        worksheet.write(row_num, 2, docx_count)

    workbook.close()
    print(f"✅ Excel file created: {output_file}")

# --- Main ---
if __name__ == "__main__":
    folder = input("Enter the folder path containing .pdf and .docx files: ")
    output_excel = "continued_counts.xlsx"

    results = match_and_count(folder)
    if results:
        write_to_excel(results, output_excel)
    else:
        print("⚠️ No matches found.")
